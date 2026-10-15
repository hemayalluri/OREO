"""
Sales reports in Excel, Word and PDF, built only from the read-only MCP sales tools.

Every number comes straight from a tool result, and the written highlights are
generated from those numbers by code, never by a language model, so a report can't
contain an invented figure. Building a report costs nothing (no API calls).
"""

from __future__ import annotations

import logging
import re
from dataclasses import dataclass, field
from datetime import date, timedelta
from pathlib import Path
from typing import Any

log = logging.getLogger("oreo.reports")

FORMATS = ("pdf", "docx", "xlsx")
# How each period sounds in a sentence: "your sales report for last month".
PERIOD_PHRASES = {
    "last_7_days": "the last 7 days", "last_30_days": "the last 30 days", "last_90_days": "the last 90 days",
    "this_month": "this month so far", "last_month": "last month", "this_quarter": "this quarter so far",
    "last_quarter": "last quarter", "this_year": "this year so far", "last_year": "last year", "all_time": "all time",
}
PERIOD_TITLES = {
    "last_7_days": "Last 7 days", "last_30_days": "Last 30 days", "last_90_days": "Last 90 days",
    "this_month": "Month to date", "last_month": "Last month", "this_quarter": "Quarter to date",
    "last_quarter": "Last quarter", "this_year": "Year to date", "last_year": "Last year",
    "all_time": "All time",
}


class ReportError(Exception):
    """A report couldn't be built; the message is safe to say out loud."""


@dataclass
class ReportData:
    title: str
    name: str             # how Oreo says it: "Europe sales report for last quarter"
    period: dict[str, str]
    filters: dict[str, str]
    summary: dict[str, Any]
    previous: dict[str, Any] | None
    change: dict[str, Any]
    top_products: list[dict[str, Any]]
    breakdown_label: str
    breakdown: list[dict[str, Any]]
    trend_label: str
    trend: list[dict[str, Any]]
    source: str
    as_of: str
    highlights: list[str] = field(default_factory=list)


@dataclass
class Report:
    data: ReportData
    files: list[Path]


# --------------------------------------------------------------------------- #
# Gathering the data
# --------------------------------------------------------------------------- #

async def gather(gateway: Any, period: str = "last_month", region: str | None = None,
                 start_date: str | None = None, end_date: str | None = None,
                 session_id: str = "") -> ReportData:
    async def call(name: str, **args: Any) -> dict[str, Any]:
        args = {k: v for k, v in args.items() if v is not None}
        res = await gateway.call(name, args, session_id)
        if res.is_error or not res.data:
            raise ReportError(_friendly_error(res.text))
        return res.data

    where = {"region": region}
    when = {"start_date": start_date, "end_date": end_date} if start_date else {"period": period}
    main = await call("get_sales_summary", **when, **where)
    p = main["period"]

    # Compare with the period of the same length just before it.
    start, end = date.fromisoformat(p["start"]), date.fromisoformat(p["end"])
    prev_end = start - timedelta(days=1)
    prev_start = prev_end - (end - start)
    previous, change = None, {}
    try:
        cmp = await call("compare_periods", a_start_date=p["start"], a_end_date=p["end"],
                         b_start_date=prev_start.isoformat(), b_end_date=prev_end.isoformat(), **where)
        if cmp["period_b"]["summary"]["orders"]:
            previous, change = {**cmp["period_b"]}, cmp["change_a_vs_b"]
    except ReportError:
        pass  # before the data starts: no comparison

    span = (end - start).days + 1
    gran = "day" if span <= 45 else "week" if span <= 200 else "month"
    dates = {"start_date": p["start"], "end_date": p["end"]}
    products = await call("top_items", dimension="product", limit=10, **dates, **where)
    by = "channel" if region else "region"
    breakdown = await call("top_items", dimension=by, limit=10, **dates, **where)
    trend = await call("sales_trend", granularity=gran, **dates, **where)

    key = p["name"].split(" ")[0]
    region_name = main["filters"].get("region")
    title = f"{region_name + ' s' if region_name else 'S'}ales report: {PERIOD_TITLES.get(key, period_text(p))}"
    when_said = PERIOD_PHRASES.get(key) or period_text(p).replace(" to ", " through ")
    name = f"{region_name + ' ' if region_name else ''}sales report for {when_said}"
    data = ReportData(
        title=title, name=name, period=p, filters=main["filters"], summary=main["summary"],
        previous=previous, change=change, top_products=products["items"],
        breakdown_label=by.capitalize(), breakdown=breakdown["items"],
        trend_label={"day": "Daily", "week": "Weekly", "month": "Monthly"}[gran], trend=trend["points"],
        source=main["source"], as_of=main["as_of"],
    )
    data.highlights = highlights(data)
    return data


def _friendly_error(text: str) -> str:
    text = re.sub(r"^Error executing tool \w+: ", "", text or "").strip()
    return text or "The sales system didn't return any data for that report."


# --------------------------------------------------------------------------- #
# Wording (by code, from the numbers)
# --------------------------------------------------------------------------- #

def money(v: float) -> str:
    v = float(v or 0)
    if abs(v) >= 1_000_000:
        return f"${v / 1_000_000:.2f}M"
    if abs(v) >= 1_000:
        return f"${v / 1_000:.1f}K"
    return f"${v:,.0f}"


def pct(v: float | None) -> str:
    return "n/a" if v is None else f"{v:+.1f}%"


def period_text(p: dict[str, str]) -> str:
    s, e = date.fromisoformat(p["start"]), date.fromisoformat(p["end"])
    return f"{s:%b %-d, %Y} to {e:%b %-d, %Y}"


def highlights(d: ReportData) -> list[str]:
    s, out = d.summary, []
    line = f"Revenue was {money(s['revenue'])} from {s['orders']:,} orders and {s['customers']:,} customers"
    rev = d.change.get("revenue", {}).get("change_pct")
    out.append(line + (f", {'up' if rev >= 0 else 'down'} {abs(rev):.1f}% on the previous period." if rev is not None else "."))
    out.append(f"Gross margin was {s['gross_margin_pct']:.1f}% ({money(s['gross_profit'])} gross profit); "
               f"the average order was {money(s['avg_order_value'])}.")
    if d.top_products:
        t = d.top_products[0]
        out.append(f"Top product: {t['name']} with {money(t['revenue'])} ({t['share_pct']:.1f}% of revenue).")
    if d.breakdown:
        b = d.breakdown[0]
        out.append(f"Biggest {d.breakdown_label.lower()}: {b['name']} with {money(b['revenue'])} ({b['share_pct']:.1f}%).")
    if d.period.get("note"):
        out.append(d.period["note"])
    return out


def spoken_money(v: float) -> str:
    """Money the way a person says it: "about 6.5 million dollars"."""
    v = float(v or 0)
    if abs(v) >= 1_000_000:
        return f"about {v / 1_000_000:.1f} million dollars".replace(".0 ", " ")
    if abs(v) >= 1_000:
        return f"about {round(v / 1_000):,} thousand dollars"
    return f"{v:,.0f} dollars"


def spoken_summary(d: ReportData) -> str:
    """One sentence Oreo can say when the report is ready."""
    rev = d.change.get("revenue", {}).get("change_pct")
    trend = f", {'up' if rev >= 0 else 'down'} {abs(rev):.0f} percent on the period before" if rev is not None else ""
    return f"Revenue came in at {spoken_money(d.summary['revenue'])}{trend}."


# --------------------------------------------------------------------------- #
# Writing the files
# --------------------------------------------------------------------------- #

def build(data: ReportData, formats: list[str], out_dir: Path) -> list[Path]:
    formats = [f for f in dict.fromkeys(formats) if f in FORMATS] or ["pdf"]
    out_dir.mkdir(parents=True, exist_ok=True)
    stem = re.sub(r"[^a-z0-9]+", "-", f"{data.title} {data.period['start']} {data.period['end']}".lower()).strip("-")
    files = []
    chart = _chart_png(data, out_dir / f"{stem}-chart.png") if {"pdf", "docx"} & set(formats) else None
    for fmt in formats:
        path = out_dir / f"{stem}.{fmt}"
        {"pdf": _write_pdf, "docx": _write_docx, "xlsx": _write_xlsx}[fmt](data, path, chart)
        files.append(path)
    if chart:
        chart.unlink(missing_ok=True)
    log.info("Built report %s", ", ".join(f.name for f in files))
    return files


def _summary_rows(d: ReportData) -> list[list[str]]:
    rows = [["Metric", "This period", "Previous period", "Change"]]
    labels = [("revenue", "Revenue", money), ("gross_profit", "Gross profit", money),
              ("gross_margin_pct", "Gross margin", lambda v: f"{v:.1f}%"), ("orders", "Orders", lambda v: f"{v:,}"),
              ("customers", "Customers", lambda v: f"{v:,}"), ("units", "Units", lambda v: f"{v:,}"),
              ("avg_order_value", "Average order", money)]
    prev = (d.previous or {}).get("summary", {})
    for key, label, fmt in labels:
        change = d.change.get(key, {}).get("change_pct")
        if key == "gross_margin_pct" and key in d.change:
            change_text = f"{d.change[key]['change']:+.1f} pts"
        else:
            change_text = pct(change) if d.previous else "n/a"
        rows.append([label, fmt(d.summary[key]), fmt(prev[key]) if key in prev else "n/a", change_text])
    return rows


def _item_rows(items: list[dict[str, Any]], label: str) -> list[list[str]]:
    return [["#", label, "Revenue", "Share"]] + [
        [str(i["rank"]), i["name"], money(i["revenue"]), f"{i['share_pct']:.1f}%"] for i in items]


def _chart_png(d: ReportData, path: Path) -> Path | None:
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except ImportError:
        return None
    fig, ax = plt.subplots(figsize=(7, 2.6), dpi=150)
    xs = [_short_label(p) for p in d.trend]
    ax.plot(range(len(xs)), [p["revenue"] for p in d.trend], color="#2a9d8f", linewidth=2)
    ax.fill_between(range(len(xs)), [p["revenue"] for p in d.trend], color="#2a9d8f", alpha=0.12)
    step = max(1, len(xs) // 8)
    ax.set_xticks(range(0, len(xs), step), [xs[i] for i in range(0, len(xs), step)], fontsize=7)
    ax.margins(x=0.01)
    ax.yaxis.set_major_formatter(lambda v, _: money(v))
    ax.tick_params(axis="y", labelsize=7)
    ax.set_title(f"{d.trend_label} revenue", fontsize=9, loc="left")
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)
    return path


def _short_label(point: dict[str, Any]) -> str:
    try:  # daily and weekly points are labelled with ISO dates; "Sep 1" fits better on an axis
        return f"{date.fromisoformat(point['label']):%b %-d}"
    except ValueError:
        return point["label"]


def _subtitle(d: ReportData) -> str:
    parts = [period_text(d.period)] + [f"{k}: {v}" for k, v in d.filters.items()]
    return " · ".join(parts)


def _footer(d: ReportData) -> str:
    return f"Source: {d.source}. Data as of {d.as_of}. Prepared by Oreo."


def _write_pdf(d: ReportData, path: Path, chart: Path | None) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(path), pagesize=letter, title=d.title, author="Oreo",
                            leftMargin=0.8 * inch, rightMargin=0.8 * inch, topMargin=0.7 * inch)
    table_style = TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#264653")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f5f4")]),
        ("ALIGN", (2, 0), (-1, -1), "RIGHT"),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4), ("TOPPADDING", (0, 0), (-1, -1), 4),
    ])
    summary_style = TableStyle(table_style.getCommands() + [("ALIGN", (1, 0), (-1, -1), "RIGHT")])
    story = [Paragraph(d.title, styles["Title"]), Paragraph(_subtitle(d), styles["Normal"]), Spacer(1, 12),
             Paragraph("Highlights", styles["Heading2"])]
    story += [Paragraph(f"• {h}", styles["Normal"]) for h in d.highlights]
    story += [Spacer(1, 10), Paragraph("Summary", styles["Heading2"]),
              Table(_summary_rows(d), style=summary_style, hAlign="LEFT")]
    if chart:
        story += [Spacer(1, 12), Image(str(chart), width=6.6 * inch, height=2.45 * inch)]
    story += [Paragraph("Top products", styles["Heading2"]),
              Table(_item_rows(d.top_products, "Product"), style=table_style, hAlign="LEFT"),
              Paragraph(f"By {d.breakdown_label.lower()}", styles["Heading2"]),
              Table(_item_rows(d.breakdown, d.breakdown_label), style=table_style, hAlign="LEFT"),
              Spacer(1, 16), Paragraph(_footer(d), styles["Italic"])]
    doc.build(story)


def _write_docx(d: ReportData, path: Path, chart: Path | None) -> None:
    from docx import Document
    from docx.shared import Inches, Pt

    doc = Document()
    doc.core_properties.title, doc.core_properties.author = d.title, "Oreo"
    doc.add_heading(d.title, level=0)
    doc.add_paragraph(_subtitle(d))
    doc.add_heading("Highlights", level=1)
    for h in d.highlights:
        doc.add_paragraph(h, style="List Bullet")

    def table(rows: list[list[str]]) -> None:
        t = doc.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Light Grid Accent 1"
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                t.cell(r, c).text = value

    doc.add_heading("Summary", level=1)
    table(_summary_rows(d))
    if chart:
        doc.add_picture(str(chart), width=Inches(6.3))
    doc.add_heading("Top products", level=1)
    table(_item_rows(d.top_products, "Product"))
    doc.add_heading(f"By {d.breakdown_label.lower()}", level=1)
    table(_item_rows(d.breakdown, d.breakdown_label))
    note = doc.add_paragraph(_footer(d))
    note.runs[0].italic, note.runs[0].font.size = True, Pt(9)
    doc.save(path)


def _write_xlsx(d: ReportData, path: Path, chart: Path | None = None) -> None:
    from openpyxl import Workbook
    from openpyxl.chart import LineChart, Reference
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    header = (Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="264653"))

    def sheet(ws, rows: list[list[Any]], widths: list[int]) -> None:
        for row in rows:
            ws.append(row)
        for cell in ws[1]:
            cell.font, cell.fill = header
        for i, w in enumerate(widths):
            ws.column_dimensions[chr(65 + i)].width = w

    ws = wb.active
    ws.title = "Summary"
    ws.append([d.title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([_subtitle(d)])
    ws.append([])
    prev = (d.previous or {}).get("summary", {})
    rows = [["Metric", "This period", "Previous period", "Change %"]]
    for key in ("revenue", "gross_profit", "gross_margin_pct", "orders", "customers", "units", "avg_order_value"):
        rows.append([key.replace("_", " ").replace("pct", "%").capitalize(), d.summary[key], prev.get(key),
                     d.change.get(key, {}).get("change_pct")])
    start = ws.max_row + 1
    for row in rows:
        ws.append(row)
    for cell in ws[start]:
        cell.font, cell.fill = header
    ws.append([])
    for h in d.highlights:
        ws.append([h])
    ws.append([])
    ws.append([_footer(d)])
    for col, w in zip("ABCD", (40, 16, 16, 12)):
        ws.column_dimensions[col].width = w

    for title, items, label in (("Top products", d.top_products, "Product"),
                                (f"By {d.breakdown_label.lower()}", d.breakdown, d.breakdown_label)):
        sheet(wb.create_sheet(title), [["Rank", label, "Revenue", "Share %"]] +
              [[i["rank"], i["name"], i["revenue"], i["share_pct"]] for i in items], [8, 32, 16, 10])

    trend = wb.create_sheet("Trend")
    sheet(trend, [["Period", "Revenue"]] + [[p["label"], p["revenue"]] for p in d.trend], [20, 16])
    line = LineChart()
    line.title, line.height, line.width = f"{d.trend_label} revenue", 7.5, 18
    line.add_data(Reference(trend, min_col=2, min_row=1, max_row=len(d.trend) + 1), titles_from_data=True)
    line.set_categories(Reference(trend, min_col=1, min_row=2, max_row=len(d.trend) + 1))
    trend.add_chart(line, "D2")
    wb.save(path)
